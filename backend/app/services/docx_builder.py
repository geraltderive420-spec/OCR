"""Génération du DOCX éditable à partir du modèle intermédiaire (PLAN.md §8).

Le builder ne dépend que de structures simples (`PageModel` / `BlockModel`) : testable sans base de données.
Structure éditable : styles Titre, paragraphes, listes, vrais tableaux Word, en-tête / pied de page réels si répétés.
Pas de zones de texte positionnées ; les pages en deux colonnes deviennent des sections Word à deux colonnes.
Taille, gras, soulignement, alignement, retraits, interligne et retours à la ligne viennent des mesures typographiques
des blocs (`typography.py`) quand elles existent ; la position d'un bloc étroit (signature à droite, date) est rendue par
ses retraits ; des blocs côte à côte (en-tête français / logo / arabe) deviennent une rangée de tableau sans bordures. Arabe : `w:bidi` sur le paragraphe, `w:rtl` sur les segments arabes,
police « complex script » définie, `w:bidiVisual` sur les tableaux majoritairement arabes.
"""
from __future__ import annotations

import html as htmlmod
import math
import re
import statistics
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html as lxml_html

from app.services.ocr.html_utils import arabic_ratio, direction_of, to_text
from app.services.typography import pt_from_cap

VISUAL = {"Picture", "Figure", "Diagram", "Image", "Stamp"}
SIZE_SCALE = [7, 8, 9, 10, 11, 12, 14, 16, 18, 20, 22, 24, 26, 28]
SERIF_FONTS = {"times new roman", "cambria", "georgia", "garamond", "book antiqua", "palatino linotype", "century"}
COL_SPACE_CM = 0.8
ARABIC_RE = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
LATIN_RE = re.compile(r"[A-Za-zÀ-ɏ]")


@dataclass
class BlockModel:
    label: str
    html: str
    bbox: tuple[float, float, float, float]
    direction: str = "ltr"
    crop_path: str | None = None  # chemin absolu d'une image
    style: dict | None = None  # mesures typographiques (typography.analyze_block)


@dataclass
class PageModel:
    width: int
    height: int
    blocks: list[BlockModel] = field(default_factory=list)
    dpi: int = 200


@dataclass
class Options:
    arabic_font: str = "Simplified Arabic"
    latin_font: str = "Arial"
    columns: bool = True  # reproduire les pages en deux colonnes

    @property
    def family(self) -> str:
        return "serif" if self.latin_font.strip().lower() in SERIF_FONTS else "sans"


# ------------------------------------------------------------------ XML utilitaires

def _rpr(run):
    return run._r.get_or_add_rPr()


def _set_fonts(run, font: str, cs_font: str | None = None) -> None:
    rpr = _rpr(run)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(attr), font)
    rf.set(qn("w:cs"), cs_font or font)


def _mark_rtl(run, size_pt: float, bold: bool) -> None:
    rpr = _rpr(run)
    rpr.append(OxmlElement("w:rtl"))
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size_pt * 2)))
    rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))


_AFTER_COLS = ("formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid",
               "printerSettings", "sectPrChange")
_AFTER_BIDI = _AFTER_COLS[_AFTER_COLS.index("bidi") + 1:]


def _insert_before(parent, el, successors: tuple[str, ...]) -> None:
    for name in successors:
        nxt = parent.find(qn(f"w:{name}"))
        if nxt is not None:
            nxt.addprevious(el)
            return
    parent.append(el)


def _bidi_paragraph(p) -> None:
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.insert(0, OxmlElement("w:bidi"))


def _classify(ch: str) -> str:
    if ARABIC_RE.match(ch):
        return "ar"
    if LATIN_RE.match(ch):
        return "la"
    return "neutral"


def _split_by_script(text: str) -> list[tuple[str, str]]:
    """[(segment, 'ar'|'la')] : les neutres (espaces, chiffres, ponctuation) suivent le segment courant."""
    out: list[list[str]] = []  # [script, texte]
    for ch in text:
        k = _classify(ch)
        if k == "neutral":
            if out:
                out[-1][1] += ch
            else:
                out.append(["la", ch])
        elif out and out[-1][0] == k:
            out[-1][1] += ch
        else:
            out.append([k, ch])
    return [(t, k) for k, t in out]


# ------------------------------------------------------------------ analyse du HTML d'un bloc

@dataclass
class Seg:
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    br: bool = False


def parse_inline(fragment: str) -> list[list[Seg]]:
    """Découpe un fragment HTML en paragraphes de segments (gras / italique / souligné / saut de ligne)."""
    fragment = fragment.strip()
    if not fragment:
        return []
    try:
        root = lxml_html.fragment_fromstring(fragment, create_parent="div")
    except Exception:  # noqa: BLE001
        return [[Seg(to_text(fragment))]]
    paras: list[list[Seg]] = [[]]

    def walk(node, bold=False, italic=False, underline=False):
        tag = node.tag if isinstance(node.tag, str) else ""
        b = bold or tag in ("b", "strong", "th")
        i = italic or tag in ("i", "em")
        u = underline or tag == "u"
        if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6") and node is not root and paras[-1]:
            paras.append([])
        if node.text:
            paras[-1].append(Seg(node.text, b, i, u))
        for child in node:
            ctag = child.tag if isinstance(child.tag, str) else ""
            if ctag == "br":
                paras[-1].append(Seg("", br=True))
            else:
                walk(child, b, i, u)
            if child.tail:
                paras[-1].append(Seg(child.tail, b, i, u))
        if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6") and node is not root:
            paras.append([])

    walk(root)
    cleaned = []
    for p in paras:
        for s in p:
            if not s.br:
                s.text = re.sub(r"\s+", " ", htmlmod.unescape(s.text))
        if any(s.br or s.text.strip() for s in p):
            cleaned.append(p)
    return cleaned


def heading_level(html: str) -> int | None:
    m = re.match(r"\s*<h([1-6])", html, re.IGNORECASE)
    return int(m.group(1)) if m else None


# ------------------------------------------------------------------ constructeur

class DocxBuilder:
    def __init__(self, pages: list[PageModel], options: Options | None = None, repeated: dict[str, str] | None = None):
        self.pages = pages
        self.opt = options or Options()
        self.repeated = repeated or {}
        self.doc = DocxDocument()
        self.page_w_cm = 21.0
        self.page_h_cm = 29.7
        self.content_w_cm = 17.0
        self.body_pt = 11
        self._env_left = self._env_right = 0.0
        self._cols = 1
        self._col_mid = 0.5
        self._target = self.doc  # conteneur courant : le document, ou une cellule de rangée
        self._in_row = False
        self._cell_w_cm = 0.0
        self._prev_bottom: float | None = None  # bas du bloc précédent (px), pour l'espace vertical

    # --- mise en page
    def _setup_section(self) -> None:
        landscape = statistics.fmean([p.width / max(1, p.height) for p in self.pages]) > 1.05
        w, h = (29.7, 21.0) if landscape else (21.0, 29.7)
        sec = self.doc.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(w), Cm(h)
        self.page_w_cm, self.page_h_cm = w, h
        lefts, rights, tops, bottoms = [], [], [], []
        for p in self.pages:
            # Un en-tête / pied qui part dans le vrai en-tête Word ne compte pas ; sinon (lettre d'une page,
            # en-tête non répété) il occupe le haut de la page et doit être dans les marges.
            moved = {k for k in ("PageHeader", "PageFooter") if self.repeated.get(k)}
            bs = [b.bbox for b in p.blocks if b.label not in moved] or [b.bbox for b in p.blocks]
            if not bs or not p.width or not p.height:
                continue
            lefts.append(min(b[0] for b in bs) / p.width)
            rights.append(1 - max(b[2] for b in bs) / p.width)
            tops.append(min(b[1] for b in bs) / p.height)
            bottoms.append(1 - max(b[3] for b in bs) / p.height)

        def cm(v: list[float], size: float, lo: float = 1.2, hi: float = 4.0) -> float:
            return min(hi, max(lo, (statistics.median(v) if v else 0.1) * size))

        sec.left_margin, sec.right_margin = Cm(cm(lefts, w)), Cm(cm(rights, w))
        sec.top_margin, sec.bottom_margin = Cm(cm(tops, h, 1.5, 3.5)), Cm(cm(bottoms, h, 1.5, 3.5))
        self.content_w_cm = w - sec.left_margin.cm - sec.right_margin.cm
        self._env_left = statistics.median(lefts) if lefts else 0.08
        self._env_right = 1 - (statistics.median(rights) if rights else 0.08)

    def _estimate_pt(self, b: BlockModel, page: PageModel) -> float | None:
        if b.style and b.style.get("cap_px"):
            # Rapport de l'arabe seulement pour un bloc vraiment arabe (une référence mixte reste en latin).
            family = "arabic" if arabic_ratio(to_text(b.html)) > 0.7 else self.opt.family
            return pt_from_cap(float(b.style["cap_px"]), page.dpi or 200, family)
        if b.style and b.style.get("pt"):
            return float(b.style["pt"])
        text = to_text(b.html)
        if len(text) < 3:
            return None
        w, h = b.bbox[2] - b.bbox[0], b.bbox[3] - b.bbox[1]
        if w <= 0 or h <= 0:
            return None
        px = min(math.sqrt(h * w / (0.6 * len(text))), h / 1.2)
        return px * 72 / (page.dpi or 200)

    @staticmethod
    def _snap(pt: float) -> int:
        return min(SIZE_SCALE, key=lambda s: abs(s - pt))

    def _body_size(self) -> int:
        """Corps du texte courant : médiane des tailles pondérée par la longueur du texte."""
        est = []
        for p in self.pages:
            for b in p.blocks:
                if b.label == "Text" and (pt := self._estimate_pt(b, p)):
                    est.extend([pt] * max(1, len(to_text(b.html)) // 50))
        return self._snap(statistics.median(est)) if est else 11

    def _block_size(self, b: BlockModel, page: PageModel) -> float:
        """Taille propre au bloc si elle s'écarte nettement du corps, sinon le corps (document homogène)."""
        pt = self._estimate_pt(b, page) if b.style else None
        if not pt:
            return self.body_pt
        snapped = self._snap(pt)
        # Mesure aberrante (signature ou tampon sur le texte) : le corps fait foi.
        if not 0.6 * self.body_pt <= snapped <= 2.0 * self.body_pt:
            return self.body_pt
        return snapped if abs(snapped - self.body_pt) >= 1 else self.body_pt

    def _alignment(self, b: BlockModel, page: PageModel):
        if not page.width:
            return None
        measured = (b.style or {}).get("align")
        if measured == "justify":
            return WD_ALIGN_PARAGRAPH.JUSTIFY
        if measured == "center":
            return WD_ALIGN_PARAGRAPH.CENTER
        if measured == "right":
            return WD_ALIGN_PARAGRAPH.RIGHT
        if self._in_row:
            return None  # la cellule donne la position
        text = to_text(b.html)
        x0, x1 = b.bbox[0] / page.width, b.bbox[2] / page.width
        env_l, env_r = self._column_env(page, b)
        env_w = max(0.1, env_r - env_l)
        width = (x1 - x0) / env_w
        center_gap = abs((x0 + x1) / 2 - (env_l + env_r) / 2)
        if len(text) < 140 and width < 0.8 and center_gap < 0.05 * env_w / 0.84:
            return WD_ALIGN_PARAGRAPH.CENTER
        if b.direction == "ltr" and (env_r - x1) < 0.03 and (x0 - env_l) > 0.35 * env_w and len(text) < 140:
            return WD_ALIGN_PARAGRAPH.RIGHT
        return None  # début de ligne : gauche (LTR) ou droite (RTL, par w:bidi)

    # --- runs
    def _column_env(self, page: PageModel, b: BlockModel) -> tuple[float, float]:
        """Bords (fractions de la largeur) de la colonne du bloc : la page entière hors section à deux colonnes."""
        if self._cols == 2 and page.width:
            if (b.bbox[0] + b.bbox[2]) / 2 / page.width < self._col_mid:
                return self._env_left, self._col_mid
            return self._col_mid, self._env_right
        return self._env_left, self._env_right

    def _add_runs(self, p, segs: list[Seg], rtl_para: bool, size: float, bold_all: bool = False,
                  underline_all: bool = False) -> None:
        for seg in segs:
            if seg.br:
                p.add_run().add_break(WD_BREAK.LINE)
                continue
            if not seg.text:
                continue
            for text, script in _split_by_script(seg.text):
                run = p.add_run(text)
                run.font.size = Pt(size)
                run.bold = seg.bold or bold_all or None
                run.italic = seg.italic or None
                run.underline = seg.underline or underline_all or None
                if script == "ar":
                    _set_fonts(run, self.opt.arabic_font, self.opt.arabic_font)
                    _mark_rtl(run, size, bool(run.bold))
                else:
                    _set_fonts(run, self.opt.latin_font, self.opt.arabic_font)

    def _paragraph(self, container, segs, direction: str, size: float, style: str | None = None,
                   align=None, bold_all: bool = False, underline_all: bool = False, fmt: dict | None = None):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        if direction == "rtl":
            _bidi_paragraph(p)
        if align is not None:
            p.alignment = align
        pf = p.paragraph_format
        pf.space_after = Pt(0) if fmt else Pt(4)
        for key, value in (fmt or {}).items():
            if value is not None:
                setattr(pf, key, value)
        self._add_runs(p, segs, direction == "rtl", size, bold_all, underline_all)
        return p

    def _block_format(self, b: BlockModel, page: PageModel, size: float, first: bool) -> dict | None:
        """Retraits (position du bloc, 1re ligne), interligne et espace avant, d'après la géométrie mesurée."""
        st = b.style
        if not st or not page.width:
            return None
        fmt: dict = {}
        dpi = page.dpi or 200
        if first and not self._in_row and self._prev_bottom is not None:
            gap_pt = max(0.0, (b.bbox[1] - self._prev_bottom) * 72 / dpi)
            fmt["space_before"] = Pt(min(72, round(gap_pt * 0.75)))  # un peu de marge : Word compose plus haut
        if st.get("pitch_px") and st.get("lines", 1) >= 2:
            ratio = st["pitch_px"] * 72 / dpi / (1.15 * size)  # Word : interligne simple ≈ 1,15 × corps
            fmt["line_spacing"] = round(min(3.0, max(1.0, ratio)) * 20) / 20
        if not self._in_row:
            env_l, env_r = self._column_env(page, b)
            env_w = max(0.1, env_r - env_l)
            col_cm = self.content_w_cm if self._cols == 1 else (self.content_w_cm - COL_SPACE_CM) / 2
            x0, x1 = b.bbox[0] / page.width, b.bbox[2] / page.width
            left = max(0.0, (x0 - env_l) / env_w * col_cm)
            right = max(0.0, (env_r - x1) / env_w * col_cm)
            if b.direction == "rtl":
                left, right = right, left  # paragraphe bidi : le retrait « gauche » de Word est à droite
            if left > 0.4:
                fmt["left_indent"] = Cm(round(left, 2))
            if right > 0.4:
                fmt["right_indent"] = Cm(round(right, 2))
        if st.get("indent"):
            width_cm = (b.bbox[2] - b.bbox[0]) / page.width * self.page_w_cm
            fmt["first_line_indent"] = Cm(round(st["indent"] * width_cm, 2))
            if st["indent"] < 0 and "left_indent" not in fmt:  # retrait suspendu : la 1re ligne dépasse à gauche
                fmt["left_indent"] = Cm(round(-st["indent"] * width_cm, 2))
            elif st["indent"] < 0:
                fmt["left_indent"] = Cm(round(fmt["left_indent"].cm - st["indent"] * width_cm, 2))
        return fmt

    @staticmethod
    def _with_breaks(b: BlockModel, paras: list[list[Seg]]) -> list[list[Seg]]:
        """Retours à la ligne des titres, adresses et signatures (le modèle les fond en une seule ligne).

        Les mots sont répartis entre les lignes au prorata de la largeur mesurée de chaque ligne.
        """
        st = b.style or {}
        widths = st.get("line_w") or []
        if (len(paras) != 1 or len(widths) < 2 or st.get("align") == "justify" or any(s.br for s in paras[0])
                or len(widths) > 6 or sum(len(s.text) for s in paras[0]) > 300):
            return paras
        words: list[tuple[str, Seg]] = []
        for seg in paras[0]:
            for w in re.findall(r"\S+\s*", seg.text):
                words.append((w, seg))
        if len(words) < len(widths):
            return paras
        total_chars = sum(len(w) for w, _ in words)
        total_w = sum(widths)
        cuts, acc, k = set(), 0, 0
        targets = [sum(widths[:i + 1]) / total_w * total_chars for i in range(len(widths) - 1)]
        for i, (w, _) in enumerate(words[:-1]):
            acc += len(w)
            if k < len(targets) and acc + len(words[i + 1][0]) / 2 >= targets[k]:
                cuts.add(i)
                k += 1
        out: list[Seg] = []
        for i, (w, seg) in enumerate(words):
            out.append(Seg(w.rstrip() if i in cuts else w, seg.bold, seg.italic, seg.underline))
            if i in cuts:
                out.append(Seg("", br=True))
        return [out]

    # --- blocs
    def _heading(self, b: BlockModel, page: PageModel, rank: dict[int, int]) -> None:
        paras = parse_inline(b.html) or [[Seg(to_text(b.html))]]
        level = heading_level(b.html) or rank.get(id(b), 2)
        level = min(max(level, 1), 3)
        measured = self._estimate_pt(b, page) if b.style else None
        size = max(self._snap(measured), self.body_pt) if measured else {1: 18, 2: 14, 3: 12}[level]
        bold = (b.style or {}).get("bold", True)
        underline = bool((b.style or {}).get("underline"))
        for k, segs in enumerate(self._with_breaks(b, paras)):
            p = self._paragraph(self._target, segs, b.direction, size, style=f"Heading {level}", align=self._alignment(b, page),
                                bold_all=bool(bold), underline_all=underline, fmt=self._block_format(b, page, size, k == 0))
            for run in p.runs:  # le style « Titre » de Word impose sa couleur : on garde l'aspect du document
                run.font.color.rgb = RGBColor(0, 0, 0)
                if not bold:
                    run.bold = False

    def _text(self, b: BlockModel, page: PageModel, style: str | None = None, size: float | None = None) -> None:
        paras = parse_inline(b.html)
        size = size or self._block_size(b, page)
        bold = bool((b.style or {}).get("bold"))
        underline = bool((b.style or {}).get("underline"))
        for k, segs in enumerate(self._with_breaks(b, paras)):
            self._paragraph(self._target, segs, b.direction, size, style=style, align=self._alignment(b, page),
                            bold_all=bold, underline_all=underline, fmt=self._block_format(b, page, size, k == 0))

    def _picture(self, b: BlockModel, page: PageModel) -> None:
        if not b.crop_path or not Path(b.crop_path).is_file():
            return
        w_rel = (b.bbox[2] - b.bbox[0]) / max(1, page.width)
        avail = self.content_w_cm if self._cols == 1 else (self.content_w_cm - COL_SPACE_CM) / 2
        if self._in_row:
            avail = self._cell_w_cm
        width = max(1.0, min(avail, w_rel * self.page_w_cm))
        para = self._target.add_paragraph()
        para.alignment = self._alignment(b, page) if not self._in_row else WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(b.crop_path, width=Cm(width))

    def _table(self, b: BlockModel, page: PageModel) -> None:
        try:
            root = lxml_html.fragment_fromstring(b.html, create_parent="div")
        except Exception:  # noqa: BLE001
            self._text(b, page)
            return
        rows = root.xpath(".//tr")
        if not rows:
            self._text(b, page)
            return
        grid: list[list[tuple | None]] = []
        ncols = 0
        for r, tr in enumerate(rows):
            while len(grid) <= r:
                grid.append([])
            c = 0
            for cell in tr.xpath("./td|./th"):
                while c < len(grid[r]) and grid[r][c] is not None:
                    c += 1
                rs = max(1, int(cell.get("rowspan", 1) or 1))
                cs = max(1, int(cell.get("colspan", 1) or 1))
                for dr in range(rs):
                    while len(grid) <= r + dr:
                        grid.append([])
                    row = grid[r + dr]
                    while len(row) < c + cs:
                        row.append(None)
                    for dc in range(cs):
                        row[c + dc] = ("span", r, c) if (dr or dc) else ("cell", cell, rs, cs)
                c += cs
            ncols = max(ncols, len(grid[r]))
        nrows = len(grid)
        if not ncols:
            return
        table = self._target.add_table(rows=nrows, cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        rtl_cells = sum(1 for tr in rows for td in tr.xpath("./td|./th") if arabic_ratio(to_text(lxml_html.tostring(td, encoding="unicode"))) > 0.5)
        total_cells = sum(len(tr.xpath("./td|./th")) for tr in rows) or 1
        if rtl_cells / total_cells > 0.5:
            tblpr = table._tbl.tblPr
            tblpr.append(OxmlElement("w:bidiVisual"))
        for r in range(nrows):
            for c in range(ncols):
                entry = grid[r][c] if c < len(grid[r]) else None
                if entry is None or entry[0] != "cell":
                    continue
                _, cell, rs, cs = entry
                tc = table.cell(r, c)
                if rs > 1 or cs > 1:
                    tc = tc.merge(table.cell(min(nrows - 1, r + rs - 1), min(ncols - 1, c + cs - 1)))
                inner = "".join(lxml_html.tostring(ch, encoding="unicode") for ch in cell) if len(cell) else ""
                frag = (cell.text or "") + inner if len(cell) else (cell.text or "")
                segs_list = parse_inline(frag) or [[Seg("")]]
                cell_dir = direction_of(to_text(frag))
                first = True
                for segs in segs_list:
                    para = tc.paragraphs[0] if first else tc.add_paragraph()
                    first = False
                    if cell_dir == "rtl":
                        _bidi_paragraph(para)
                    self._add_runs(para, segs, cell_dir == "rtl", max(8, self.body_pt - 1), bold_all=cell.tag == "th")
        if not self._in_row:
            self._target.add_paragraph()

    def _set_header_footer(self) -> None:
        sec = self.doc.sections[0]
        for label, part in (("PageHeader", sec.header), ("PageFooter", sec.footer)):
            html = self.repeated.get(label)
            if not html:
                continue
            para = part.paragraphs[0]
            segs_list = parse_inline(html) or [[Seg(to_text(html))]]
            direction = direction_of(to_text(html))
            for k, segs in enumerate(segs_list):
                target = para if k == 0 else part.add_paragraph()
                if direction == "rtl":
                    _bidi_paragraph(target)
                self._add_runs(target, segs, direction == "rtl", 9)

    # --- orchestration
    def build(self) -> DocxDocument:
        if not self.pages:
            return self.doc
        self._setup_section()
        self.body_pt = self._body_size()
        rank = self._heading_ranks()
        self._set_header_footer()
        repeated_keys = {k: re.sub(r"\d+", "#", to_text(v).lower()).strip() for k, v in self.repeated.items()}
        rtl_doc = self._mostly_rtl()
        first = True
        for page in self.pages:
            for cols, group in self._column_groups(page) or [(1, [])]:  # une page vide reste une page
                starts_page = not group or not group[0][1]
                if first:
                    self._set_columns(self.doc.sections[0], cols, rtl_doc)
                    first = False
                elif cols != self._cols or starts_page:
                    start = WD_SECTION.NEW_PAGE if starts_page else WD_SECTION.CONTINUOUS
                    self._set_columns(self.doc.add_section(start), cols, rtl_doc)
                self._cols = cols
                if starts_page:
                    self._prev_bottom = None
                blocks = [b for b, _ in group]
                for row in (self._rows(blocks) if cols == 1 else [[b] for b in blocks]):
                    if len(row) == 1:
                        self._block(row[0], page, rank, repeated_keys)
                    else:
                        self._row(row, page, rank)
                    self._prev_bottom = max(b.bbox[3] for b in row)
        return self.doc

    @staticmethod
    def _rows(blocks: list[BlockModel]) -> list[list[BlockModel]]:
        """Regroupe les blocs consécutifs posés côte à côte (même hauteur, sans recouvrement horizontal)."""
        rows: list[list[BlockModel]] = []
        for b in blocks:
            row = rows[-1] if rows else None
            if row and b.label != "Table" and all(o.label != "Table" for o in row):
                h = min(b.bbox[3] - b.bbox[1], min(o.bbox[3] - o.bbox[1] for o in row))
                v_overlap = min(b.bbox[3], max(o.bbox[3] for o in row)) - max(b.bbox[1], min(o.bbox[1] for o in row))
                side_by_side = all(b.bbox[0] >= o.bbox[2] - 5 or b.bbox[2] <= o.bbox[0] + 5 for o in row)
                if side_by_side and h > 0 and v_overlap >= 0.5 * h:
                    row.append(b)
                    continue
            rows.append([b])
        return rows

    def _row(self, row: list[BlockModel], page: PageModel, rank: dict[int, int]) -> None:
        """Blocs côte à côte : une rangée de tableau sans bordures, une cellule par bloc, de gauche à droite."""
        row = sorted(row, key=lambda b: b.bbox[0])
        env_l, env_r = self._env_left * page.width, self._env_right * page.width
        bounds = [min(env_l, row[0].bbox[0])]
        for a, c in zip(row, row[1:]):
            bounds.append((a.bbox[2] + c.bbox[0]) / 2)
        bounds.append(max(env_r, row[-1].bbox[2]))
        span = bounds[-1] - bounds[0]
        widths = [(bounds[i + 1] - bounds[i]) / span * self.content_w_cm for i in range(len(row))]
        if self._prev_bottom is not None:  # espace vertical avant la rangée
            gap_pt = max(0.0, (min(b.bbox[1] for b in row) - self._prev_bottom) * 72 / (page.dpi or 200))
            if gap_pt > 6:
                spacer = self._target.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(min(72, round(gap_pt * 0.75)))
        table = self._target.add_table(rows=1, cols=len(row))
        table.autofit = False  # largeurs fixes (python-docx place w:tblLayout au bon endroit du schéma)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w * 567)))  # la grille fait foi pour Word, pas seulement les cellules
        for i, (b, w) in enumerate(zip(row, widths)):
            cell = table.cell(0, i)
            cell.width = Cm(w)
            self._target, self._in_row, self._cell_w_cm = cell, True, w
            try:
                self._render(b, page, rank)
                for p in cell.paragraphs:  # cellule du milieu (logo, titre) : centrée
                    if p.alignment is None and 0 < i < len(row) - 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first = cell.paragraphs[0]
                if not first.text and not first._p.xpath(".//w:drawing") and len(cell.paragraphs) > 1:
                    first._p.getparent().remove(first._p)  # paragraphe vide créé avec la cellule
            finally:
                self._target, self._in_row = self.doc, False

    def _heading_ranks(self) -> dict[int, int]:
        """Niveau 1 à 3 des titres d'après leur taille mesurée (à défaut, la hauteur d'une ligne de la boîte)."""
        sizes: dict[int, float] = {}
        for p in self.pages:
            for b in p.blocks:
                if b.label == "SectionHeader":
                    pt = self._estimate_pt(b, p) or (b.bbox[3] - b.bbox[1]) * 72 / (p.dpi or 200)
                    sizes[id(b)] = round(pt, 1)
        distinct = sorted(set(sizes.values()), reverse=True)
        rank = {}
        for k, pt in sizes.items():
            pos = distinct.index(pt) / max(1, len(distinct) - 1) if len(distinct) > 1 else 0.5
            rank[k] = 1 if pos < 0.25 else (2 if pos < 0.65 else 3)
        return rank

    def _mostly_rtl(self) -> bool:
        n = sum(len(to_text(b.html)) for p in self.pages for b in p.blocks if b.direction == "rtl")
        total = sum(len(to_text(b.html)) for p in self.pages for b in p.blocks) or 1
        return n / total > 0.5

    def _set_columns(self, section, cols: int, rtl: bool) -> None:
        sectpr = section._sectPr
        for old in sectpr.findall(qn("w:cols")):
            sectpr.remove(old)
        el = OxmlElement("w:cols")
        el.set(qn("w:num"), str(cols))
        el.set(qn("w:space"), str(int(COL_SPACE_CM * 567)))
        _insert_before(sectpr, el, _AFTER_COLS)  # l'ordre des enfants de sectPr est imposé par le schéma
        bidi = sectpr.find(qn("w:bidi"))
        if rtl and cols > 1 and bidi is None:  # colonnes de droite à gauche pour un document arabe
            _insert_before(sectpr, OxmlElement("w:bidi"), _AFTER_BIDI)
        elif bidi is not None and not (rtl and cols > 1):
            sectpr.remove(bidi)

    def _column_groups(self, page: PageModel) -> list[tuple[int, list[tuple[BlockModel, bool]]]]:
        """Découpe la page en suites de blocs à 1 ou 2 colonnes, dans l'ordre de lecture.

        Chaque bloc est accompagné d'un drapeau « pas le premier de la page » (False = la page commence).
        Une suite de blocs étroits (moins de 55 % de la largeur utile) répartis des deux côtés du milieu
        devient une section à deux colonnes ; les blocs pleine largeur (titres, tableaux larges) coupent la suite.
        """
        blocks = list(page.blocks)
        if not blocks:
            return []
        main = [b for b in blocks if b.label not in ("PageHeader", "PageFooter")] or blocks
        left = min(b.bbox[0] for b in main)
        right = max(b.bbox[2] for b in main)
        span = max(1.0, right - left)
        mid = left + span / 2
        if page.width:
            self._col_mid = mid / page.width

        slack = 0.04 * span

        def side(b: BlockModel) -> str | None:
            """Colonne du bloc, ou None s'il chevauche le milieu (bloc centré : titre court, signature…)."""
            if b.bbox[2] <= mid + slack:
                return "g"
            if b.bbox[0] >= mid - slack:
                return "d"
            return None

        def narrow(b: BlockModel) -> bool:
            return ((b.bbox[2] - b.bbox[0]) < 0.55 * span and side(b) is not None
                    and b.label not in ("PageHeader", "PageFooter"))

        runs: list[tuple[bool, list[BlockModel]]] = []
        for b in blocks:
            n = narrow(b) and self.opt.columns
            if runs and runs[-1][0] == n:
                runs[-1][1].append(b)
            else:
                runs.append((n, [b]))
        groups: list[tuple[int, list[BlockModel]]] = []
        for n, bs in runs:
            two = n and len({side(b) for b in bs}) == 2 and sum(len(to_text(b.html)) for b in bs) > 200
            cols = 2 if two else 1
            if groups and groups[-1][0] == cols:
                groups[-1][1].extend(bs)
            else:
                groups.append((cols, bs))
        out = []
        first = True
        for cols, bs in groups:
            out.append((cols, [(b, not first or k > 0) for k, b in enumerate(bs)]))
            first = False
        return out

    def _block(self, b: BlockModel, page: PageModel, rank: dict[int, int], repeated_keys: dict[str, str]) -> None:
        key = re.sub(r"\d+", "#", to_text(b.html).lower()).strip()
        if b.label in ("PageHeader", "PageFooter") and repeated_keys.get(b.label) == key and key:
            return  # déjà dans le vrai en-tête / pied de page
        self._render(b, page, rank)

    def _render(self, b: BlockModel, page: PageModel, rank: dict[int, int]) -> None:
        if b.label in VISUAL:
            self._picture(b, page)
        elif b.label == "Table":
            self._table(b, page)
        elif b.label == "SectionHeader":
            self._heading(b, page, rank)
        elif b.label in ("ListGroup", "ListItem"):
            self._text(b, page, style="List Bullet")
        elif b.label == "Caption":
            self._text(b, page, style="Caption", size=max(8, self.body_pt - 2))
        elif b.label == "Footnote":
            self._text(b, page, size=max(8, self.body_pt - 3))
        else:
            self._text(b, page)

    def save(self, path: str | Path) -> None:
        self.build().save(str(path))
