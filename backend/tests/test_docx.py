from docx import Document

from app.services.docx_builder import BlockModel, DocxBuilder, PageModel, parse_inline, _split_by_script


def make_pages():
    w, h = 1654, 2339
    return [
        PageModel(w, h, [
            BlockModel("PageHeader", "Ministère de test", (200, 40, 1400, 110)),
            BlockModel("SectionHeader", "<h1>Titre principal</h1>", (300, 200, 1350, 300)),
            BlockModel("Text", "Un paragraphe en <b>gras</b> et normal, référence 2024/118.", (200, 330, 1450, 520)),
            BlockModel("Text", "هذا نص عربي تجريبي رقم 45 مع كلمات", (200, 540, 1450, 640), direction="rtl"),
            BlockModel("Text", "Mixte : voir المادة 12 du décret", (200, 660, 1450, 740)),
            BlockModel("Table", "<table><tr><th>Nom</th><th>Qté</th><th>Prix</th></tr>"
                                "<tr><td>Alpha</td><td>3</td><td>10</td></tr>"
                                "<tr><td colspan='2'>Total</td><td>30</td></tr>"
                                "<tr><td rowspan='2'>Beta</td><td>1</td><td>5</td></tr>"
                                "<tr><td>2</td><td>6</td></tr></table>", (200, 800, 1450, 1200)),
            BlockModel("ListGroup", "<ul><li>Premier</li><li>Second</li></ul>", (200, 1250, 1000, 1400)),
        ]),
        PageModel(w, h, [
            BlockModel("PageHeader", "Ministère de test", (200, 40, 1400, 110)),
            BlockModel("Text", "Deuxième page.", (200, 200, 1450, 260)),
            BlockModel("PageFooter", "Page 2", (700, 2250, 950, 2300)),
        ]),
    ]


def build(tmp_path, **kw):
    out = tmp_path / "out.docx"
    DocxBuilder(make_pages(), repeated={"PageHeader": "Ministère de test"}, **kw).save(out)
    return Document(str(out))


def test_structure_de_base(tmp_path):
    d = build(tmp_path)
    styles = [p.style.name for p in d.paragraphs if p.text.strip()]
    assert "Heading 1" in styles
    assert "List Bullet" in styles
    assert len(d.tables) == 1
    # l'en-tête répété est un vrai en-tête Word, pas un paragraphe du corps
    assert "Ministère de test" in d.sections[0].header.paragraphs[0].text
    assert not any("Ministère de test" in p.text for p in d.paragraphs)
    # saut de page entre les pages du scan (saut de section « nouvelle page »)
    from docx.enum.section import WD_SECTION
    assert len(d.sections) == 2 and d.sections[1].start_type == WD_SECTION.NEW_PAGE


def test_tableau_fusions(tmp_path):
    t = build(tmp_path).tables[0]
    assert len(t.rows) == 5 and len(t.columns) == 3
    xml = t._tbl.xml
    assert xml.count("w:gridSpan") == 1  # « Total » sur 2 colonnes
    assert "w:vMerge" in xml  # « Beta » sur 2 lignes
    assert t.cell(0, 0).paragraphs[0].runs[0].bold


def test_arabe_bidi_et_segments(tmp_path):
    d = build(tmp_path)
    ar = next(p for p in d.paragraphs if "عربي" in p.text)
    assert "<w:bidi" in ar._p.xml
    assert any("<w:rtl" in r._r.xml for r in ar.runs)
    mixed = next(p for p in d.paragraphs if p.text.startswith("Mixte"))
    assert "<w:bidi" not in mixed._p.xml  # paragraphe de base LTR
    kinds = [("<w:rtl" in r._r.xml) for r in mixed.runs]
    assert any(kinds) and not all(kinds)  # segments arabes et latins séparés


def test_decoupage_par_ecriture():
    segs = _split_by_script("Voir المادة 12 du décret")
    assert [k for _, k in segs] == ["la", "ar", "la"]
    assert "".join(s for s, _ in segs) == "Voir المادة 12 du décret"


def test_gras_inline():
    paras = parse_inline("Un <b>mot</b> clé<br>ligne 2")
    flat = [s for p in paras for s in p]
    assert any(s.bold and s.text == "mot" for s in flat)
    assert any(s.br for s in flat)


def test_page_vide_et_orientation_paysage(tmp_path):
    out = tmp_path / "l.docx"
    DocxBuilder([PageModel(2339, 1654, [BlockModel("Text", "Paysage", (100, 100, 2000, 200))])]).save(out)
    sec = Document(str(out)).sections[0]
    assert sec.page_width > sec.page_height
    DocxBuilder([]).save(tmp_path / "v.docx")  # aucun plantage sans page


def _two_column_page(rtl=False):
    w, h = 1654, 2339
    txt = ("نص عربي طويل للتجربة " if rtl else "Texte de colonne assez long pour compter. ") * 8
    d = "rtl" if rtl else "ltr"
    left = [BlockModel("Text", txt, (150, 400 + 300 * k, 800, 650 + 300 * k), direction=d) for k in range(3)]
    right = [BlockModel("Text", txt, (860, 400 + 300 * k, 1500, 650 + 300 * k), direction=d) for k in range(3)]
    return PageModel(w, h, [
        BlockModel("SectionHeader", "Titre pleine largeur", (150, 200, 1500, 300), direction=d),
        *(right + left if rtl else left + right),
        BlockModel("Text", "Signature pleine largeur " * 6, (150, 1400, 1500, 1500), direction=d),
    ])


def test_deux_colonnes_en_section_word(tmp_path):
    out = tmp_path / "c.docx"
    DocxBuilder([_two_column_page()]).save(out)
    d = Document(str(out))
    cols = [s._sectPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cols") for s in d.sections]
    nums = [c.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num") for c in cols]
    assert nums == ["1", "2", "1"]  # titre, colonnes, signature


def test_colonnes_arabes_de_droite_a_gauche(tmp_path):
    out = tmp_path / "a.docx"
    DocxBuilder([_two_column_page(rtl=True)]).save(out)
    d = Document(str(out))
    assert "<w:bidi/>" in d.sections[1]._sectPr.xml


def test_colonnes_desactivees(tmp_path):
    from app.services.docx_builder import Options
    out = tmp_path / "u.docx"
    DocxBuilder([_two_column_page()], Options(columns=False)).save(out)
    assert len(Document(str(out)).sections) == 1


def test_mesures_typographiques_appliquees(tmp_path):
    w, h = 1654, 2339
    pages = [PageModel(w, h, [
        BlockModel("SectionHeader", "Un titre", (200, 100, 1400, 180), style={"pt": 15.6, "bold": True, "align": "center"}),
        BlockModel("Text", "Corps du texte. " * 20, (200, 200, 1450, 600), style={"pt": 10.2, "align": "justify"}),
        BlockModel("Text", "Note en petit. " * 5, (200, 700, 1450, 760), style={"pt": 7.1}),
    ])]
    out = tmp_path / "t.docx"
    DocxBuilder(pages).save(out)
    d = Document(str(out))
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    title, body, note = [p for p in d.paragraphs if p.text.strip()]
    assert title.runs[0].font.size.pt == 16 and title.runs[0].bold and title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert body.runs[0].font.size.pt == 10 and body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert note.runs[0].font.size.pt == 7
    assert body.runs[0]._r.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Arial"


def test_blocs_etroits_centres_sans_colonnes(tmp_path):
    w, h = 1654, 2339
    txt = "Ligne centrée assez longue pour dépasser le seuil de texte. " * 3
    page = PageModel(w, h, [BlockModel("Text", txt, (500, 300 + 150 * k, 1150, 420 + 150 * k)) for k in range(4)]
                     + [BlockModel("Text", "Pleine largeur " * 10, (150, 1200, 1500, 1300))])
    out = tmp_path / "n.docx"
    DocxBuilder([page]).save(out)
    assert len(Document(str(out)).sections) == 1


def test_taille_aberrante_ramenee_au_corps(tmp_path):
    w, h = 1654, 2339
    pages = [PageModel(w, h, [
        BlockModel("Text", "Corps du texte. " * 30, (200, 200, 1450, 600), style={"pt": 10.1}),
        BlockModel("Text", "Nom sous une signature", (200, 700, 900, 900), style={"pt": 27.0}),
    ])]
    out = tmp_path / "s.docx"
    DocxBuilder(pages).save(out)
    sizes = [p.runs[0].font.size.pt for p in Document(str(out)).paragraphs if p.text.strip()]
    assert sizes == [10, 10]


def test_blocs_cote_a_cote_en_rangee_et_mise_en_forme_mesuree(tmp_path):
    """Lettre : en-tête français / arabe côte à côte, titre sur deux lignes, signature à droite soulignée."""
    w, h = 2339, 1654
    page = PageModel(w, h, [
        BlockModel("PageHeader", "ETABLISSEMENT DE LA PLATEFORME", (430, 30, 1080, 100), style={"cap_px": 21, "lines": 1}),
        BlockModel("PageHeader", "مؤسسة قاعدة المنظومات", (1570, 30, 1930, 100), direction="rtl", style={"cap_px": 23, "lines": 1}),
        BlockModel("SectionHeader", "MONSIEUR, LE DIRECTEUR DU DPR-SEF", (890, 420, 1460, 510),
                   style={"cap_px": 30, "lines": 2, "align": "center", "bold": True, "line_w": [0.38, 0.94]}),
        BlockModel("Text", "Suite au disfonctionnement des groupes. " * 8, (423, 764, 1932, 1006),
                   style={"cap_px": 26, "lines": 5, "align": "justify", "indent": 0.09, "pitch_px": 48}),
        BlockModel("Text", "LE DIRECTEUR GENERAL DE L'EPIC-PSE/2°RM", (1347, 1072, 1870, 1169),
                   style={"cap_px": 29, "lines": 2, "align": "center", "bold": True, "underline": True, "line_w": [0.93, 0.74]}),
    ])
    out = tmp_path / "lettre.docx"
    DocxBuilder([page]).save(out)
    d = Document(str(out))
    assert len(d.tables) == 1 and len(d.tables[0].columns) == 2  # en-tête côte à côte
    assert "ETABLISSEMENT" in d.tables[0].cell(0, 0).text and "مؤسسة" in d.tables[0].cell(0, 1).text
    paras = [p for p in d.paragraphs if p.text.strip()]
    title, body, sign = paras
    assert "<w:br/>" in title._p.xml and title.runs[0].bold  # deux lignes
    assert body.paragraph_format.first_line_indent.cm > 0.5 and body.paragraph_format.line_spacing >= 1.2
    assert sign.runs[0].underline and sign.paragraph_format.left_indent.cm > 8  # à droite de la page
    assert sign.runs[0].font.size.pt == 14 and body.runs[0].font.size.pt == 12
